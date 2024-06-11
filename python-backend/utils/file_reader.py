import PyPDF2
import pandas as pd
from openpyxl import load_workbook
from docx import Document


def read_pdf(file_path):
    try:
        with open(file_path, "rb") as file:
            reader = PyPDF2.PdfReader(file)
            text = ""
            for page in reader.pages:
                text += page.extract_text()
            return text
    except Exception as e:
        print(f"Error reading PDF: {e}")
        return None


def read_xlsx(file_path):
    try:
        workbook = load_workbook(file_path)
        sheet = workbook.active
        data = sheet.values
        return "\n".join([", ".join(map(str, row)) for row in data])
    except Exception as e:
        print(f"Error reading XLSX: {e}")
        return None


def read_csv(file_path):
    try:
        df = pd.read_csv(file_path)
        return df.to_csv(index=False)
    except Exception as e:
        print(f"Error reading CSV: {e}")
        return None


def read_docx(file_path):
    try:
        doc = Document(file_path)
        return "\n".join([para.text for para in doc.paragraphs])
    except Exception as e:
        print(f"Error reading DOCX: {e}")
        return None


def read_file(file_path):
    if file_path.endswith(".pdf"):
        return read_pdf(file_path)
    elif file_path.endswith(".xlsx"):
        return read_xlsx(file_path)
    elif file_path.endswith(".csv"):
        return read_csv(file_path)
    elif file_path.endswith(".docx"):
        return read_docx(file_path)
    else:
        print(f"Unsupported file format: {file_path}")
        return None
